from pathlib import Path
import asyncio
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database.db import get_excel_employee_rows, get_total_strikes, get_setting

EXCEL_PATH = Path(__file__).resolve().parent / "employee_records.xlsx"
_excel_lock = asyncio.Lock()

STATUS_LABELS = {
    "active": "على رأس العمل",
    "vacation": "إجازة",
    "fired": "مفصول",
    "resigned": "مستقيل",
}

async def sync_employee_excel(guild):
    async with _excel_lock:
        """ينشئ/يحدّث ملف Excel من قاعدة البيانات. قاعدة البيانات هي المصدر الأساسي."""
        rows = await get_excel_employee_rows(guild.id)
        target_raw = await get_setting(guild.id, "weekly_invoice_target")
        try:
            target = max(1, int(target_raw or 15))
        except ValueError:
            target = 15
    
        wb = Workbook()
        ws = wb.active
        ws.title = "الموظفين"
        ws.sheet_view.rightToLeft = True
    
        headers = [
            "Discord ID", "اسم الموظف", "رقم الموظف", "Citizen ID", "تاريخ التوظيف",
            "الحالة", "إعادة التقديم", "إجمالي السترايكات",
            "إجمالي الفواتير", "إجمالي ساعات العمل", "إجمالي المهام", "إجمالي النقاط",
            "فواتير الأسبوع", "المطلوب أسبوعيًا", "حالة التفاعل", "ساعات الأسبوع",
            "مهام الأسبوع", "نقاط الأسبوع", "أيام الإجازة", "بداية الإجازة", "نهاية الإجازة",
        ]
        ws.append(headers)
    
        for row in rows:
            (uid, game_name, employee_no, citizen_id, hired_at, status, reapply_allowed,
             total_invoices, total_work_seconds, total_tasks, total_points,
             weekly_invoices, weekly_work_seconds, weekly_tasks, weekly_points,
             vacation_days, vacation_start, vacation_end) = row
            member = guild.get_member(int(uid))
            if game_name and not str(game_name).startswith("Discord ") and str(game_name) != "-":
                display_name = game_name
            else:
                display_name = member.display_name if member else (game_name or str(uid))
            strikes = await get_total_strikes(guild.id, int(uid))
            interaction_status = "مستثنى - إجازة" if status == "vacation" else ("متفاعل" if weekly_invoices >= target else "غير متفاعل")
            ws.append([
                str(uid), display_name, employee_no, citizen_id, hired_at,
                STATUS_LABELS.get(status, status), "مسموح" if reapply_allowed else "غير مسموح", strikes,
                total_invoices, round((total_work_seconds or 0) / 3600, 2), total_tasks, total_points,
                weekly_invoices, target, interaction_status, round((weekly_work_seconds or 0) / 3600, 2),
                weekly_tasks, weekly_points, vacation_days or "", vacation_start or "", vacation_end or "",
            ])
    
        header_fill = PatternFill("solid", fgColor="2F241D")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D9D9D9")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
    
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
    
        widths = [22, 24, 16, 18, 16, 18, 16, 18, 16, 18, 14, 14, 16, 16, 18, 16, 14, 14, 14, 22, 22]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    
        wb.save(EXCEL_PATH)
        return EXCEL_PATH
